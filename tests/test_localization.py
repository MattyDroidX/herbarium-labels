"""Language selection contracts, using synthetic readings and injected clients only."""
from collections import Counter
from contextlib import redirect_stdout
from dataclasses import replace
import hashlib
from importlib.resources import files
import io
import json
import os
from pathlib import Path
from string import Formatter
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from pydantic import BaseModel

from herbarium_labels import (
    TranscriptionOptions, TranscriptionProfile, get_profile, load_messages,
    read_image, transcribe_folder,
)
from herbarium_labels import cli, pipeline, sennen
from herbarium_labels.cache import identity
from herbarium_labels.prompts import effective_prompt
from herbarium_labels.report_messages import DEFAULT_REPORT_MESSAGES
from herbarium_labels.sennen_suggestions import DEFAULT_MESSAGES
from test_reuse import photo, provider, sheet


def reading_for(profile):
    if profile.name == 'sennen':
        return sheet(uncertainties=['Cota: leitura incerta.'],
                     labels=[sennen.Label(origin='Sennen', literal_text='Plantes d’Espagne\n= N.º 12')])
    values = {key: '' for key in profile.schema.model_fields}
    values.update(rotation_degrees=0, literal_text='= N.º 12\nPlantes d’Espagne',
                  name_as_written='Sphagnum palustre L.', handwritten_notes='Rev. à vérifier',
                  uncertainties=['Nome: leitura incerta.'])
    return profile.schema(**values)


class NotesLanguageTests(unittest.TestCase):
    def test_none_matches_frozen_a3_prompt_bytes_and_cache_identity(self):
        # Captured from unmodified 0.1.0a3 before implementing language options.
        frozen = json.loads((Path(__file__).parent / 'fixtures/a3-contracts.json').read_text(encoding='utf-8'))
        with tempfile.TemporaryDirectory() as folder:
            image = Path(folder) / 'image'
            image.write_bytes(b'a3 cache identity fixture')
            for name, contract in frozen.items():
                profile = get_profile(name)
                for options in (TranscriptionOptions(model='offline-model'),
                                TranscriptionOptions(model='offline-model', notes_language=None)):
                    with self.subTest(profile=name, options=options):
                        client = provider(reading_for(profile))
                        read_image(photo(), profile=name, client=client, options=options)
                        prompt = client.messages.parse.call_args.kwargs['messages'][0]['content'][-1]['text']
                        self.assertEqual(prompt.encode('utf-8'), profile.prompt.encode('utf-8'))
                        self.assertEqual(hashlib.sha256(prompt.encode('utf-8')).hexdigest(), contract['prompt_sha256'])
                        config = pipeline.cache_configuration(profile, options, version_fn=lambda package: 'baseline-' + package)
                        self.assertEqual(identity(image, config), contract['cache'])

    def test_invalid_notes_languages_raise_value_error_at_construction(self):
        for language in ('', 'PT', 'pt-PT', 'fr', 'en-US', 0, True, [], {}):
            with self.subTest(language=language), self.assertRaises(ValueError):
                TranscriptionOptions(notes_language=language)

    def test_language_rule_is_shared_and_does_not_rewrite_literal_values(self):
        for name in ('po-bryophytes', 'sennen'):
            profile = get_profile(name)
            original_prompt = profile.prompt
            reading = reading_for(profile)
            before = reading.model_dump_json()
            for language, display in (('en', 'English'), ('pt', 'Portuguese'), ('es', 'Spanish')):
                with self.subTest(profile=name, language=language):
                    client = provider(reading)
                    actual, _ = read_image(photo(), profile=name, client=client,
                                          options=TranscriptionOptions(model='offline-model', notes_language=language))
                    prompt = client.messages.parse.call_args.kwargs['messages'][0]['content'][-1]['text']
                    self.assertEqual(prompt, effective_prompt(profile, language))
                    self.assertTrue(prompt.startswith(original_prompt + '\n\n'))
                    rule = prompt[len(original_prompt):]
                    self.assertIn('only the entries of `uncertainties` in ' + display, rule)
                    self.assertIn('Do not translate or normalize', rule)
                    for field in ('literal_text', 'labels', 'name_as_written', 'handwritten_notes'):
                        self.assertIn('`' + field + '`', rule)
                    self.assertEqual(actual.model_dump_json(), before)
                    self.assertEqual(profile.prompt, original_prompt)
                    client.close.assert_not_called()

    def test_external_profile_gets_the_same_language_rule(self):
        class ExternalReading(BaseModel):
            literal_text: str
            handwritten_notes: str
            uncertainties: list[str]
        profile = TranscriptionProfile('external-language-test', 'Original external contract.\n', ExternalReading,
                                       (('filename', 'Filename'), ('handwritten_notes', 'Notes')))
        reading = ExternalReading(literal_text='Fleurs des Pyrénées', handwritten_notes='à revoir', uncertainties=['Dúvida.'])
        client = provider(reading)
        actual, _ = read_image(photo(), profile=profile, client=client,
                               options=TranscriptionOptions(notes_language='pt'))
        prompt = client.messages.parse.call_args.kwargs['messages'][0]['content'][-1]['text']
        self.assertEqual(prompt, effective_prompt(profile, 'pt'))
        self.assertEqual(actual, reading)

    def test_cache_changes_for_each_language_and_contains_the_sent_prompt(self):
        for name in ('sennen', 'po-bryophytes'):
            configs = []
            for language in (None, 'en', 'pt', 'es'):
                options = TranscriptionOptions(model='offline-model', notes_language=language)
                profile = get_profile(name)
                config = pipeline.cache_configuration(profile, options)
                self.assertEqual(config['prompt'], effective_prompt(profile, language))
                configs.append(json.dumps(config, sort_keys=True))
            self.assertEqual(len(set(configs)), 4)

    def test_json_records_language_and_resume_rejects_a_language_change_before_requests(self):
        with tempfile.TemporaryDirectory() as folder, redirect_stdout(io.StringIO()):
            images = Path(folder) / 'input'
            images.mkdir()
            (images / 'sheet.jpg').write_bytes(photo())
            output = Path(folder) / 'output'
            client = provider()
            options = TranscriptionOptions(model='offline-model', notes_language='pt', workers=1)
            transcribe_folder(images, output, profile='sennen', client=client, options=options)
            target = output / 'json/sheet.json'
            original = target.read_bytes()
            data = json.loads(original)
            self.assertEqual(data['notes_language'], 'pt')
            self.assertNotIn('notes_language', data['reading'])
            self.assertEqual(transcribe_folder(images, output, profile='sennen', client=client, options=options).reused, 1)
            with self.assertRaisesRegex(ValueError, 'incompatible'):
                transcribe_folder(images, output, profile='sennen', client=client,
                                  options=replace(options, notes_language='es'))
            self.assertEqual(client.messages.parse.call_count, 1)
            self.assertEqual(target.read_bytes(), original)

    def test_a3_reading_without_language_metadata_reuses_unchanged(self):
        with tempfile.TemporaryDirectory() as folder, redirect_stdout(io.StringIO()):
            images = Path(folder)
            image = images / 'sheet.jpg'
            image.write_bytes(photo())
            output = images / 'output'
            (output / 'json').mkdir(parents=True)
            options = TranscriptionOptions(model='offline-model', workers=1)
            old_config = pipeline.cache_configuration('sennen', options)
            old_config['prompt'] = get_profile('sennen').prompt
            record = {'schema_version': '0.2.0', 'filename': image.name, 'profile': 'sennen',
                      'model': options.model, 'input_tokens': 10, 'output_tokens': 20,
                      'cache': identity(image, old_config), 'reading': sheet().model_dump()}
            target = output / 'json/sheet.json'
            target.write_text(json.dumps(record), encoding='utf-8')
            before = target.read_bytes()
            with patch.dict(os.environ, {}, clear=True):
                result = transcribe_folder(images, output, profile='sennen', options=options)
            self.assertEqual(result.reused, 1)
            self.assertEqual(target.read_bytes(), before)


class CatalogueTests(unittest.TestCase):
    def test_catalogues_have_exactly_matching_keys_and_format_markers(self):
        def flatten(value, prefix=()):
            result = {}
            for key, item in value.items():
                if isinstance(item, dict):
                    result.update(flatten(item, prefix + (key,)))
                else:
                    self.assertIsInstance(item, str)
                    self.assertTrue(item.strip())
                    result[prefix + (key,)] = item
            return result
        def markers(text):
            return Counter((key, spec, conversion) for _, key, spec, conversion in Formatter().parse(text)
                           if key is not None)
        root = files('herbarium_labels').joinpath('locales')
        english = flatten(json.loads(root.joinpath('en.json').read_text(encoding='utf-8')))
        for lang in ('pt', 'es'):
            translated = flatten(json.loads(root.joinpath(lang + '.json').read_text(encoding='utf-8')))
            self.assertEqual(english.keys(), translated.keys(), lang)
            for key, default in english.items():
                with self.subTest(lang=lang, key=key):
                    self.assertEqual(markers(default), markers(translated[key]))

    def test_english_catalogue_matches_all_existing_defaults(self):
        for name in ('po-bryophytes', 'sennen'):
            expected = {**DEFAULT_REPORT_MESSAGES,
                        **{f'field.{key}': label for key, label in get_profile(name).columns if key is not None}}
            self.assertEqual(load_messages('en', profile=name), expected)
        self.assertEqual(load_messages('en', kind='review'), dict(DEFAULT_MESSAGES))
        self.assertEqual(load_messages('en'), load_messages('en', profile='po-bryophytes'))

    def test_portuguese_uses_existing_report_terminology(self):
        root = Path(__file__).resolve().parents[1] / 'examples/synthetic'
        for profile, suffix in (('po-bryophytes', 'po'), ('sennen', 'sennen')):
            example = json.loads((root / f'report-messages-{suffix}-pt.json').read_text(encoding='utf-8'))
            actual = load_messages('pt', profile=profile)
            for key, text in example.items():
                if key != 'pdf.title':  # Bundled titles retain the original {profile} marker.
                    self.assertEqual(actual[key], text)

    def test_load_messages_rejects_invalid_language_kind_and_review_profile(self):
        for lang in ('PT', 'pt-PT', 'fr', None, [], 1):
            with self.subTest(lang=lang), self.assertRaises(ValueError):
                load_messages(lang)
        with self.assertRaises(ValueError):
            load_messages('en', kind='other')
        with self.assertRaises(ValueError):
            load_messages('pt', kind='review', profile='po-bryophytes')

    def test_loaded_mappings_do_not_leak_mutations_between_calls(self):
        for kind in ('report', 'review'):
            original = load_messages('pt', kind=kind)
            changed = load_messages('pt', kind=kind)
            changed[next(iter(changed))] = 'HOST OVERRIDE'
            self.assertEqual(load_messages('pt', kind=kind), original)

    def test_external_profile_uses_common_labels_and_its_own_columns(self):
        profile = replace(get_profile('sennen'), name='external-report')
        text = load_messages('pt', profile=profile)
        self.assertEqual(text['column.filename'], 'Ficheiro')
        self.assertEqual(text['field.filename'], 'Ficheiro')
        self.assertEqual(text['field.shelfmark'], 'Shelfmark')

    def test_portuguese_review_preserves_taxa_legacy_wording(self):
        messages = load_messages('pt', kind='review')
        self.assertEqual({key: messages[key] for key in ('source', 'exsiccata_number')}, {
            'source': 'Sugestão automática ({model}) a partir da fotografia {filename}. Confirme na imagem antes de aceitar.',
            'exsiccata_number': 'exsiccata_number atual: «{current}». Leitura na etiqueta Sennen: «{reading}».',
        })
        reading = sheet(sennen_exsiccata_number='N.º 12',
                        uncertainties=['exsiccata_number: dúvida original.'])
        rows = sennen.suggestions({'transcription': 'Reviewed', 'exsiccata_number': '7'},
                                  reading, 'sheet.jpg', messages=messages, model_label='mock-model')
        self.assertEqual(rows, [('exsiccata_number', '12',
            'Sugestão automática (mock-model) a partir da fotografia sheet.jpg. '
            'Confirme na imagem antes de aceitar. '
            'exsiccata_number atual: «7». Leitura na etiqueta Sennen: «N.º 12».'
            '\n\nDúvidas assinaladas pelo modelo:\n- exsiccata_number: dúvida original.')])

    def test_sennen_review_catalogues_render_with_host_fields_and_preserve_literal_doubts(self):
        reading = sheet(shelfmark='68 / 9', uncertainties=['Cota: dúvida original.', 'shelfmark: source doubt.'])
        for lang in ('en', 'pt', 'es'):
            with self.subTest(lang=lang):
                result = sennen.suggestions({'transcription': 'Reviewed', 'cota': '12'}, reading, 'sheet.jpg',
                    messages=load_messages(lang, kind='review'), field_map={'shelfmark': 'cota'},
                    options=TranscriptionOptions(model='offline-model', notes_language='pt'))
                self.assertEqual(result[0][:2], ('cota', '68/9'))
                self.assertIn('offline-model', result[0][2])
                self.assertIn('Cota: dúvida original.', result[0][2])
                self.assertIn('shelfmark: source doubt.', result[0][2])
                self.assertEqual(reading.shelfmark, '68 / 9')


class LanguageCLITests(unittest.TestCase):
    def test_notes_language_reaches_transcribe_and_batch_options_independently_of_lang(self):
        with tempfile.TemporaryDirectory() as folder:
            for command in ('transcribe', 'batch'):
                argv = ['herbarium-labels', command, '--images', folder, '--output', folder + '/out',
                        '--notes-language', 'pt', '--profile', 'sennen']
                if command == 'batch':
                    argv += ['--lang', 'es']
                with patch.object(sys, 'argv', argv), patch.object(pipeline, command + '_folder') as run:
                    run.return_value.failures = 0
                    cli.main()
                self.assertEqual(run.call_args.kwargs['options'].notes_language, 'pt')
                if command == 'batch':
                    self.assertEqual(run.call_args.kwargs['messages'], load_messages('es', profile='sennen'))

    def test_report_lang_pt_generates_portuguese_workbook_and_json_overrides_merge_key_by_key(self):
        with tempfile.TemporaryDirectory() as folder, redirect_stdout(io.StringIO()):
            images = Path(folder) / 'input'
            images.mkdir()
            (images / 'sheet.jpg').write_bytes(photo())
            output = Path(folder) / 'out'
            transcribe_folder(images, output, profile='sennen', client=provider(),
                              options=TranscriptionOptions(model='offline-model', workers=1))
            original = (output / 'json/sheet.json').read_bytes()
            overrides = Path(folder) / 'overrides.json'
            overrides.write_text(json.dumps({'column.filename': 'Ficheiro do Taxa'}), encoding='utf-8')
            argv = ['herbarium-labels', 'report', '--profile', 'sennen', '--images', str(images),
                    '--output', str(output), '--lang', 'pt', '--messages-json', str(overrides), '--no-pdf']
            with patch.object(sys, 'argv', argv):
                cli.main()
            book = load_workbook(output / 'transcriptions.xlsx')
            try:
                self.assertEqual(book.sheetnames, ['Transcrição literal', 'Campos'])
                self.assertEqual(book.worksheets[0]['A1'].value, 'Ficheiro do Taxa')
                self.assertEqual(book.worksheets[0]['E1'].value, 'Dúvidas de leitura')
                self.assertEqual(book.worksheets[0]['F1'].value, 'Revisto')
                self.assertIn('Cota', [cell.value for cell in book.worksheets[1][1]])
                self.assertIn('= SYNTHETIC A', book.worksheets[0]['C2'].value)
            finally:
                book.close()
            self.assertEqual((output / 'json/sheet.json').read_bytes(), original)

    def test_batch_lang_alone_leaves_notes_language_unset(self):
        with tempfile.TemporaryDirectory() as folder:
            argv = ['herbarium-labels', 'batch', '--images', folder, '--lang', 'pt']
            with patch.object(sys, 'argv', argv), patch.object(pipeline, 'batch_folder') as run:
                cli.main()
            self.assertIsNone(run.call_args.kwargs['options'].notes_language)
            self.assertEqual(run.call_args.kwargs['messages']['sheet.literal'], 'Transcrição literal')

    def test_invalid_cli_language_codes_are_rejected(self):
        for command, flag in (('transcribe', '--notes-language'), ('batch', '--notes-language'),
                              ('report', '--lang'), ('batch', '--lang')):
            argv = ['herbarium-labels', command, '--images', '.', '--output', './out', flag, 'fr']
            with self.subTest(command=command, flag=flag), patch.object(sys, 'argv', argv), redirect_stdout(io.StringIO()):
                with patch('sys.stderr', new=io.StringIO()), self.assertRaises(SystemExit) as error:
                    cli.main()
                self.assertEqual(error.exception.code, 2)


if __name__ == '__main__':
    unittest.main()
