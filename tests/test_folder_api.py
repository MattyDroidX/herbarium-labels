"""Public folder APIs, plugin profiles, report language, and explicit legacy reuse."""
from contextlib import contextmanager, redirect_stdout
from dataclasses import replace
import hashlib
from importlib.metadata import EntryPoint
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import Mock, patch

from PIL import Image
from pydantic import BaseModel

from herbarium_labels import (
    TokenPrices, TranscriptionOptions, TranscriptionProfile, batch_folder, create_report,
    get_profile, list_profiles, register_profile, transcribe_folder,
)
from herbarium_labels import cli, pipeline, profiles, sennen
from herbarium_labels.pricing import cost_text, estimate_cost
from herbarium_labels.provider import DEFAULT_MAX_RETRIES
from herbarium_labels.report_messages import DEFAULT_REPORT_MESSAGES
from test_reuse import photo, provider, sheet


def old_record(image, reading=None, *, cached=True):
    return {'schema_version': '0.2.0', 'filename': image.name, 'profile': 'sennen',
            'model': 'historical-model', 'input_tokens': 30, 'output_tokens': 40,
            **({'cache': {'version': 1, 'image_sha256': hashlib.sha256(image.read_bytes()).hexdigest(),
                          'configuration_sha256': '0' * 64}} if cached else {}),
            'reading': (reading or sheet()).model_dump()}


@contextmanager
def legacy_folder(*, cached=True):
    with tempfile.TemporaryDirectory() as folder:
        images = Path(folder) / 'images'
        images.mkdir()
        image = images / 'old.jpg'
        image.write_bytes(photo())
        output = images / 'transcription'
        (output / 'json').mkdir(parents=True)
        target = output / 'json/old.json'
        target.write_text(json.dumps(old_record(image, cached=cached), ensure_ascii=False))
        yield images, output, image, target


class FolderAPITests(unittest.TestCase):
    def test_plain_paths_work_for_all_folder_apis_and_return_artifacts(self):
        with tempfile.TemporaryDirectory() as folder:
            images = Path(folder) / 'images'
            images.mkdir()
            (images / 'sheet.jpg').write_bytes(photo())
            output = Path(folder) / 'delivery'
            client = provider()
            totals = transcribe_folder(str(images), output, profile='sennen', client=client, limit=1)
            self.assertEqual(totals.photos, 1)
            report = create_report(images, output, profile='sennen', include_pdf=False)
            self.assertTrue(report.workbook.is_file())
            self.assertIsNone(report.pdf)
            batch = batch_folder(images, output, profile='sennen', client=client)
            self.assertEqual(batch.totals.reused, 1)
            self.assertTrue(batch.archive.is_file())
            self.assertEqual(batch.archive.name, 'delivery.zip')
            self.assertEqual(client.messages.parse.call_count, 1)
            client.close.assert_not_called()

    def test_folder_api_applies_host_request_context(self):
        calls = []
        @contextmanager
        def limit():
            calls.append('enter')
            try:
                yield
            finally:
                calls.append('exit')
        with tempfile.TemporaryDirectory() as folder:
            images = Path(folder)
            (images / 'sheet.jpg').write_bytes(photo())
            transcribe_folder(images, images / 'out', profile='sennen', client=provider(), request_context=limit)
        self.assertEqual(calls, ['enter', 'exit'])

    def test_invalid_limit_never_calls_provider(self):
        with tempfile.TemporaryDirectory() as folder:
            client = provider()
            for limit in (0, -1, 1.5):
                with self.subTest(limit=limit), self.assertRaises(ValueError):
                    transcribe_folder(folder, Path(folder) / 'out', client=client, limit=limit)
            client.messages.parse.assert_not_called()

    def test_legacy_batch_configuration_failure_remains_an_exception(self):
        from herbarium_labels import TranscriptionError
        with tempfile.TemporaryDirectory() as folder, patch.dict(os.environ, {}, clear=True):
            root = Path(folder)
            (root / 'sheet.jpg').write_bytes(photo())
            with self.assertRaises(TranscriptionError):
                pipeline.batch(SimpleNamespace(images=root, pdf=False), profile='sennen')

    def test_cli_uses_plain_parameters_and_shared_retry_default(self):
        with tempfile.TemporaryDirectory() as folder:
            argv = ['herbarium-labels', 'transcribe', '--images', folder, '--output', folder + '/out']
            with patch.object(sys, 'argv', argv), patch.object(pipeline, 'transcribe_folder',
                    return_value=SimpleNamespace(failures=0)) as run:
                cli.main()
            self.assertEqual(run.call_args.args, (folder, folder + '/out'))
            self.assertEqual(run.call_args.kwargs['options'].max_retries, DEFAULT_MAX_RETRIES)
            self.assertEqual(TranscriptionOptions().max_retries, DEFAULT_MAX_RETRIES)

    def test_explicit_prices_override_environment_and_model_labels_use_options(self):
        prices = TokenPrices(1, 2, model='explicit-model')
        with patch.dict(os.environ, {'TRANSCRIPTION_PRICE_IN': '999', 'TRANSCRIPTION_PRICE_OUT': '999',
                                    'TRANSCRIPTION_MODEL': 'environment-model'}):
            self.assertEqual(estimate_cost(1_000_000, 2_000_000, prices=prices, model='explicit-model'), 5)
            self.assertIsNone(estimate_cost(1, 1, prices=prices, model='different-model'))
            self.assertEqual(cost_text(1_000_000, 0, prices=TokenPrices(0, 0)), '0.0000 USD')
            options = TranscriptionOptions(model='explicit-model', prices=prices)
            rows = sennen.suggestions({}, sheet(), 'sheet.jpg', options=options)
            self.assertIn('(explicit-model)', rows[0][2])
            self.assertNotIn('environment-model', rows[0][2])
            self.assertEqual(sennen.MODEL, 'environment-model')
            self.assertEqual(sennen.estimated_cost('explicit-model', 1_000_000, 0, prices=prices), 1)
        with self.assertRaises(ValueError):
            TokenPrices(float('nan'), 1)

    def test_batch_prices_override_options_for_both_stages_without_mutating_options(self):
        options = TranscriptionOptions(model='unit', prices=TokenPrices(1_000_000, 1_000_000, model='unit'))
        override = TokenPrices(2_000_000, 3_000_000, model='unit')
        with tempfile.TemporaryDirectory() as folder, redirect_stdout(io.StringIO()) as stream:
            images = Path(folder)
            (images / 'sheet.jpg').write_bytes(photo())
            batch_folder(images, profile='sennen', client=provider(), options=options, prices=override)
        self.assertEqual(stream.getvalue().count('80.0000 USD'), 2)
        self.assertEqual(options.prices, TokenPrices(1_000_000, 1_000_000, model='unit'))

    def test_explicit_prices_reach_transcription_summary(self):
        with tempfile.TemporaryDirectory() as folder, redirect_stdout(io.StringIO()) as stream:
            images = Path(folder)
            (images / 'sheet.jpg').write_bytes(photo())
            transcribe_folder(images, images / 'out', profile='sennen', client=provider(),
                options=TranscriptionOptions(model='unit', prices=TokenPrices(1_000_000, 1_000_000, model='unit')))
        self.assertIn('30.0000 USD', stream.getvalue())


class ProfileExtensionTests(unittest.TestCase):
    def profile(self, name='synthetic-extension'):
        return replace(get_profile('sennen'), name=name)

    def test_register_profile_is_available_to_cli_and_rejects_collisions(self):
        with patch.dict(profiles._REGISTRY, dict(profiles._REGISTRY), clear=True), \
                patch.object(profiles, 'entry_points', return_value=[]):
            custom = self.profile()
            register_profile(custom)
            self.assertIs(get_profile(custom.name), custom)
            self.assertIn(custom.name, list_profiles())
            with self.assertRaises(ValueError):
                register_profile(custom)
            with patch.object(sys, 'argv', ['herbarium-labels', 'profiles']), redirect_stdout(io.StringIO()) as output:
                cli.main()
            self.assertIn(custom.name, output.getvalue())

    def test_plugins_are_lazy_and_factories_must_match_the_entry_name(self):
        custom = self.profile()
        loader = Mock(return_value=lambda: custom)
        ep = SimpleNamespace(name=custom.name, value='extension:factory', load=loader)
        with patch.object(profiles, 'entry_points', return_value=[ep]):
            self.assertIn(custom.name, list_profiles())
            loader.assert_not_called()
            self.assertEqual(get_profile(custom.name), custom)
        wrong = SimpleNamespace(name='wrong-name', value='extension:factory', load=Mock(return_value=lambda: custom))
        with patch.object(profiles, 'entry_points', return_value=[wrong]), self.assertRaises(ValueError):
            get_profile('wrong-name')

    def test_duplicate_plugin_names_cannot_silently_override_a_profile(self):
        for name in ('sennen', 'duplicate-plugin'):
            ep = SimpleNamespace(name=name, value='extension:factory', load=Mock())
            entries = [ep] if name == 'sennen' else [ep, ep]
            with patch.object(profiles, 'entry_points', return_value=entries), self.assertRaises(ValueError):
                list_profiles()
            ep.load.assert_not_called()

    def test_installed_entry_point_works_in_a_fresh_cli_process(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / 'synthetic_profile.py').write_text('''from herbarium_labels import TranscriptionProfile\nfrom pydantic import BaseModel\nclass Reading(BaseModel):\n    literal_text: str\n    uncertainties: list[str]\ndef create_profile():\n    return TranscriptionProfile("fixture-herbarium", "Copy text.", Reading, (("filename", "File"),))\n''')
            metadata = root / 'synthetic_profile-1.0.dist-info'
            metadata.mkdir()
            (metadata / 'METADATA').write_text('Name: synthetic-profile\nVersion: 1.0\n')
            (metadata / 'entry_points.txt').write_text('[herbarium_labels.profiles]\nfixture-herbarium = synthetic_profile:create_profile\n')
            images = root / 'images'
            images.mkdir()
            (images / 'sheet.jpg').write_bytes(photo())
            output = root / 'output'
            (output / 'json').mkdir(parents=True)
            (output / 'json/sheet.json').write_text(json.dumps({'profile': 'fixture-herbarium',
                'filename': 'sheet.jpg', 'model': 'offline', 'input_tokens': 0, 'output_tokens': 0,
                'reading': {'literal_text': 'Synthetic text', 'uncertainties': []}}))
            environment = dict(os.environ)
            environment['PYTHONPATH'] = str(root) + os.pathsep + environment.get('PYTHONPATH', '')
            command = [sys.executable, '-m', 'herbarium_labels', 'report', '--profile', 'fixture-herbarium',
                       '--images', str(images), '--output', str(output), '--no-pdf']
            result = subprocess.run(command, env=environment, capture_output=True, text=True, timeout=30)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertTrue((output / 'transcriptions.xlsx').is_file())


class ExistingReadingTests(unittest.TestCase):
    def test_explicit_acceptance_preserves_json_and_only_calls_api_for_new_photos(self):
        with legacy_folder() as (images, output, image, target):
            before = target.read_bytes()
            (images / 'new.jpg').write_bytes(photo((110, 200)))
            client = provider()
            with self.assertRaisesRegex(ValueError, 'new output directory'):
                transcribe_folder(images, output, profile='sennen', client=client)
            client.messages.parse.assert_not_called()
            result = batch_folder(images, profile='sennen', client=client, accept_existing=True)
            self.assertEqual((result.totals.photos, result.totals.reused), (1, 1))
            self.assertEqual(target.read_bytes(), before)
            ledger = json.loads((output / '.accepted-existing.json').read_text())
            self.assertTrue(ledger['records']['old.json']['historical_source_verified'])
            self.assertEqual(json.loads(target.read_text())['model'], 'historical-model')
            self.assertEqual(client.messages.parse.call_count, 1)
            with patch.dict(os.environ, {}, clear=True):
                resumed = transcribe_folder(images, output, profile='sennen')
            self.assertEqual(resumed.reused, 2)

    def test_hashless_legacy_reading_is_explicitly_marked_unverified(self):
        with legacy_folder(cached=False) as (images, output, image, target):
            before = target.read_bytes()
            client = provider()
            self.assertEqual(transcribe_folder(images, output, profile='sennen', client=client,
                                               accept_existing=True).reused, 1)
            client.messages.parse.assert_not_called()
            ledger = json.loads((output / '.accepted-existing.json').read_text())
            self.assertFalse(ledger['records']['old.json']['historical_source_verified'])
            self.assertEqual(target.read_bytes(), before)
            create_report(images, output, profile='sennen', include_pdf=False)

    def test_acceptance_cannot_bypass_a_recorded_or_accepted_image_hash(self):
        with legacy_folder() as (images, output, image, target):
            before = target.read_bytes()
            image.write_bytes(photo((120, 200)))
            client = provider()
            with self.assertRaisesRegex(ValueError, 'image differs'):
                transcribe_folder(images, output, profile='sennen', client=client, accept_existing=True)
            self.assertEqual(target.read_bytes(), before)
            self.assertFalse((output / '.accepted-existing.json').exists())
            client.messages.parse.assert_not_called()
        with legacy_folder(cached=False) as (images, output, image, target):
            transcribe_folder(images, output, profile='sennen', accept_existing=True)
            image.write_bytes(photo((130, 200)))
            with self.assertRaises(ValueError):
                transcribe_folder(images, output, profile='sennen', accept_existing=True)
            with self.assertRaises(ValueError):
                create_report(images, output, profile='sennen', include_pdf=False)

    def test_all_old_records_are_checked_before_acceptance_is_committed(self):
        with legacy_folder() as (images, output, image, target):
            (images / 'zzz.jpg').write_bytes(photo())
            (output / 'json/zzz.json').write_text('{broken')
            client = provider()
            with self.assertRaises(ValueError):
                transcribe_folder(images, output, profile='sennen', client=client, accept_existing=True)
            self.assertFalse((output / '.accepted-existing.json').exists())
            client.messages.parse.assert_not_called()

    def test_post_acceptance_human_edits_need_explicit_renewal_without_overwrite(self):
        with legacy_folder() as (images, output, image, target):
            transcribe_folder(images, output, profile='sennen', accept_existing=True)
            item = json.loads(target.read_text())
            item['reading']['labels'][0]['literal_text'] = 'HUMAN REVIEW\nPreserved'
            target.write_text(json.dumps(item))
            before = target.read_bytes()
            with self.assertRaisesRegex(ValueError, 'changed'):
                transcribe_folder(images, output, profile='sennen')
            transcribe_folder(images, output, profile='sennen', accept_existing=True)
            self.assertEqual(target.read_bytes(), before)
            entry = json.loads((output / '.accepted-existing.json').read_text())['records']['old.json']
            self.assertIsNotNone(entry['previous_acceptance'])

    def test_old_portuguese_envelope_is_reused_without_rewriting_keys(self):
        from test_english_api import legacy_reading
        with legacy_folder(cached=False) as (images, output, image, target):
            item = {'ficheiro': image.name, 'modelo': 'historical', 'input_tokens': 1,
                    'output_tokens': 2, 'leitura': legacy_reading()}
            target.write_text(json.dumps(item, ensure_ascii=False))
            before = target.read_bytes()
            self.assertEqual(transcribe_folder(images, output, profile='po-bryophytes', accept_existing=True).reused, 1)
            self.assertEqual(target.read_bytes(), before)
            create_report(images, output, profile='po-bryophytes', include_pdf=False)

    def test_cli_accepts_existing_and_uses_json_report_messages_without_a_key(self):
        from openpyxl import load_workbook
        with legacy_folder() as (images, output, image, target):
            mapping = images.parent / 'messages.json'
            mapping.write_text(json.dumps({'sheet.literal': 'Leitura', 'sheet.fields': 'Campos'}))
            argv = ['herbarium-labels', 'batch', '--profile', 'sennen', '--images', str(images),
                    '--accept-existing', '--messages-json', str(mapping)]
            with patch.object(sys, 'argv', argv), patch.dict(os.environ, {}, clear=True):
                cli.main()
            book = load_workbook(output / 'transcriptions.xlsx')
            try:
                self.assertEqual(book.sheetnames, ['Leitura', 'Campos'])
            finally:
                book.close()
            self.assertTrue((output / '.accepted-existing.json').is_file())

    def test_invalid_reading_or_wrong_profile_cannot_be_accepted(self):
        with legacy_folder() as (images, output, image, target):
            client = provider()
            with self.assertRaises(ValueError):
                transcribe_folder(images, output, profile='po-bryophytes', client=client, accept_existing=True)
            item = json.loads(target.read_text())
            item['input_tokens'] = -1
            target.write_text(json.dumps(item))
            with self.assertRaises(ValueError):
                transcribe_folder(images, output, profile='sennen', client=client, accept_existing=True)
            client.messages.parse.assert_not_called()


class ReportLanguageTests(unittest.TestCase):
    def test_mapping_translates_workbook_fields_and_pdf_metadata_without_touching_literal(self):
        from openpyxl import load_workbook
        messages = {'sheet.literal': 'Transcrição literal', 'sheet.fields': 'Campos',
                    'column.literal_text': 'Transcrição literal', 'column.reviewed': 'Revisto',
                    'column.handwritten_notes': 'Notas manuscritas', 'column.uncertainties': 'Dúvidas de leitura',
                    'pdf.title': 'Revisão de transcrição', 'review.choices': '[ ] sim    [ ] corrigido',
                    'field.shelfmark': 'Cota'}
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / 'sheet.jpg').write_bytes(photo())
            output = root / 'out'
            transcribe_folder(root, output, profile='sennen', client=provider())
            before = (output / 'json/sheet.json').read_bytes()
            files = create_report(root, output, profile='sennen', messages=messages)
            book = load_workbook(files.workbook)
            try:
                self.assertEqual(book.sheetnames, ['Transcrição literal', 'Campos'])
                self.assertEqual(book.worksheets[0]['F1'].value, 'Revisto')
                self.assertIn('Cota', [cell.value for cell in book.worksheets[1][1]])
                self.assertIn('= SYNTHETIC A', book.worksheets[0]['C2'].value)
                self.assertEqual(book.worksheets[0]['C2'].data_type, 's')
            finally:
                book.close()
            self.assertTrue(files.pdf.read_bytes().startswith(b'%PDF'))
            self.assertEqual((output / 'json/sheet.json').read_bytes(), before)

    def test_callback_resolves_once_and_invalid_sheet_names_stop_before_requests(self):
        calls = []
        def translate(key, context):
            calls.append(key)
            return {'sheet.literal': 'Leitura', 'sheet.fields': 'Campos'}.get(key, context['default']).format_map(context)
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / 'sheet.jpg').write_bytes(photo())
            batch_folder(root, profile='sennen', client=provider(), include_pdf=True, messages=translate)
            self.assertEqual(len(calls), len(set(calls)))
        for messages in ({'sheet.literal': 'bad/name'}, {'sheet.fields': 'Literal transcription'},
                         {'sheet.literal': 'a' * 32}, {'typo': 'x'}):
            with tempfile.TemporaryDirectory() as folder:
                client = provider()
                with self.subTest(messages=messages), self.assertRaises(ValueError):
                    batch_folder(folder, profile='sennen', client=client, messages=messages)
                client.messages.parse.assert_not_called()

    def test_sennen_suggestion_module_keeps_old_imports_and_three_host_scenarios(self):
        from herbarium_labels.sennen_suggestions import DEFAULT_MESSAGES
        from herbarium_labels.suggestions import DEFAULT_MESSAGES as compatibility
        self.assertIs(DEFAULT_MESSAGES, compatibility)
        for record, reading in [({'cota': '68 / 9', 'transcription': 'reviewed'}, sheet(shelfmark='68/9')),
                                ({'collection_date': '1993-06-15'}, sheet(collection_date='15-VI-1993')),
                                ({'cota': '12'}, sheet(shelfmark='68 / 9', uncertainties=['Cota ilegível no círculo']))]:
            before = dict(record)
            rows = sennen.suggestions(record, reading, 'sheet.jpg', field_map={'shelfmark': 'cota'})
            self.assertEqual(record, before)
            if 'collection_date' in record:
                self.assertNotIn('collection_date', [key for key, _, _ in rows])
            elif record.get('cota') == '68 / 9':
                self.assertNotIn('cota', [key for key, _, _ in rows])
            else:
                self.assertIn('Cota ilegível no círculo', next(text for key, _, text in rows if key == 'cota'))


if __name__ == '__main__':
    unittest.main()
