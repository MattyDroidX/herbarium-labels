"""Profile-driven cache, transcription, and offline Excel/PDF/ZIP delivery."""
import io
import json
import time
import zipfile
from concurrent.futures import ThreadPoolExecutor
from dataclasses import asdict, dataclass, replace
from importlib.metadata import PackageNotFoundError, version
from pathlib import Path
from xml.sax.saxutils import escape

from PIL import Image
from .pricing import cost_text, estimate_cost
from .cache import TaskResult, digest, identity, output_lock
from .existing import AcceptedReadings
from .report_messages import report_messages
from .errors import TranscriptionError
from .imaging import upright, prepare_images, register_heif_opener
from .legacy import RECORD_SCHEMA_VERSION, normalize_record
from .profiles import get_profile
from .provider import TranscriptionOptions, create_client, close_client, request_reading

def safe_photo(folder, name):
    root = Path(folder).resolve()
    candidate = (root / name).resolve()
    if Path(name).name != name or not candidate.is_relative_to(root) or not candidate.is_file():
        raise ValueError('Image reference is outside the input directory or does not exist.')
    return candidate



def photos(folder):
    files = sorted(p for p in Path(folder).iterdir() if p.is_file() and p.suffix.lower() in ('.jpg', '.jpeg', '.png', '.heic', '.heif'))
    stems = [p.stem.casefold() for p in files]
    if len(stems) != len(set(stems)):
        raise ValueError('Duplicate image filename stems: rename the images before processing.')
    for path in files:
        safe_photo(folder, path.name)
    return files


def cache_configuration(profile, options, *, version_fn=version):
    profile = get_profile(profile)
    try:
        provider_version = version_fn('anthropic')
    except PackageNotFoundError:
        provider_version = None
    return {'profile': profile.name, 'model': options.resolved_model,
            'prompt': profile.prompt, 'schema': profile.schema.model_json_schema(),
            'anthropic': provider_version, 'pydantic': version_fn('pydantic'),
            'max_tokens': options.max_tokens, 'transform': 'shared-exif-rgb-thumbnail-jpeg-v2',
            'images': asdict(profile.images), 'pillow': version_fn('Pillow'),
            'pillow_heif': version_fn('pillow-heif') if register_heif_opener else None}


class _IncompleteBatchError(TranscriptionError):
    """Preserve legacy exit behavior only for incomplete delivery."""


@dataclass(frozen=True)
class ReportFiles:
    workbook: Path
    pdf: Path | None


@dataclass(frozen=True)
class BatchResult:
    totals: TaskResult
    reports: ReportFiles
    archive: Path


def transcribe_folder(images, output, *, profile='po-bryophytes', options=None, client=None,
                      limit=None, sample=False, accept_existing=False, request_context=None):
    """Transcribe missing images; retain existing readings and return task totals.

    Cost estimates use options.prices, then environment rates when omitted.
    """
    profile, options = get_profile(profile), options or TranscriptionOptions()
    with output_lock(output):
        return _transcribe(Path(images), Path(output), profile, options, client,
                           limit=limit, sample=sample, accept_existing=accept_existing,
                           request_context=request_context)


def transcribe(args, *, profile='po-bryophytes', options=None, client=None, _configuration=None):
    """Compatibility adapter for the earlier argparse-shaped API."""
    profile, options = get_profile(profile), options or TranscriptionOptions()
    with output_lock(args.output):
        return _transcribe(Path(args.images), Path(args.output), profile, options, client,
                           configuration=_configuration, limit=getattr(args, 'limit', None),
                           sample=getattr(args, 'sample', False),
                           accept_existing=getattr(args, 'accept_existing', False))


def _transcribe(images, output, profile, options, client=None, configuration=None, *,
                limit=None, sample=False, accept_existing=False, request_context=None):
    if limit is not None and (type(limit) is not int or limit <= 0):
        raise ValueError('limit must be a positive integer.')
    out = output / 'json'
    out.mkdir(parents=True, exist_ok=True)
    files = photos(images)
    if sample and limit:
        step = max(1, len(files) // limit)
        files = files[::step][:limit]
    elif limit:
        files = files[:limit]
    # Check the entire selected batch before making any paid request.
    configuration = configuration or cache_configuration(profile, options)
    tasks = []
    accepted = AcceptedReadings(output)
    for path in files:
        target = out / (path.stem + '.json')
        expected = identity(path, configuration)
        cached = accepted.reusable(target, path, expected, profile, accept_existing=accept_existing)
        tasks.append((path, target, expected, cached, digest(target.read_bytes()) if cached else None))
    accepted.commit()
    owned = client is None
    if owned and any(not cached for _, _, _, cached, _ in tasks):
        client = create_client(options)

    def run(task):
        path, target, expected, cached, prior_hash = task
        if cached:
            if digest(path.read_bytes()) != expected['image_sha256'] or digest(target.read_bytes()) != prior_hash:
                raise ValueError('An existing reading/image changed after preflight; run again after review.')
            return TaskResult(reused=1)
        content = path.read_bytes()
        if digest(content) != expected['image_sha256']:
            raise ValueError(f'{path.name}: image changed during the run; run the command again.')
        started = time.time()
        try:
            blocks = prepare_images(content, profile.images)
            reading, usage = request_reading(client, blocks, profile, options, request_context)
        except (OSError, Image.DecompressionBombError):
            print(f'{path.name}: image preparation failed')
            return TaskResult(failures=1)
        except TranscriptionError as error:
            print(f'{path.name}: {error}')
            return TaskResult(failures=1, input_tokens=error.input_tokens,
                              output_tokens=error.output_tokens)
        partial = target.with_suffix('.json.partial')
        partial.write_text(json.dumps({'schema_version':RECORD_SCHEMA_VERSION, 'filename':path.name, 'profile':profile.name, 'model':options.resolved_model, 'elapsed_seconds':round(time.time() - started, 1),
                                      'input_tokens':usage['input_tokens'], 'output_tokens':usage['output_tokens'],
                                      'cache':expected, 'reading':reading.model_dump()}, ensure_ascii=False, indent=2), encoding='utf-8')
        if target.exists():
            raise ValueError(f'{target.name}: a result appeared during the run; it was not replaced.')
        partial.replace(target)
        return TaskResult(photos=1, input_tokens=usage['input_tokens'],
                          output_tokens=usage['output_tokens'])

    try:
        with ThreadPoolExecutor(max_workers=options.workers) as pool:
            results = list(pool.map(run, tasks))
    finally:
        if owned:
            close_client(client)
    totals = TaskResult(**{name: sum(getattr(result, name) for result in results)
                           for name in TaskResult.__dataclass_fields__})
    cost = cost_text(totals.input_tokens, totals.output_tokens, prices=options.prices, model=options.resolved_model)
    print(f'transcribed {totals.photos} · reused {totals.reused} · failures {totals.failures} '
          f'· tokens {totals.input_tokens}+{totals.output_tokens} · estimated cost: {cost}')
    return totals


def read_records(output, *, profile='po-bryophytes'):
    profile = get_profile(profile)
    accepted = AcceptedReadings(output)
    results = []
    for path in sorted((Path(output) / 'json').glob('*.json')):
        item = normalize_record(json.loads(path.read_text(encoding='utf-8')))
        entry = accepted.validate_report(path)
        recorded_profile = item.get('profile', entry['profile'] if entry else 'po-bryophytes')
        if recorded_profile != profile.name:
            raise ValueError('Stored reading profile differs; choose the matching profile.')
        item['reading'] = profile.schema.model_validate(item['reading']).model_dump()
        results.append(item)
    return results


def report_photo(folder, item):
    path = safe_photo(folder, item['filename'])
    recorded_hash = (item.get('cache') or {}).get('image_sha256')
    if recorded_hash is not None and digest(path.read_bytes()) != recorded_hash:
        raise ValueError(f"{path.name}: image differs from the transcribed source; restore it or use a new output directory.")
    return path


def create_report(images, output, *, profile='po-bryophytes', include_pdf=True, messages=None, prices=None):
    """Build offline reports; explicit prices override the environment fallback."""
    profile = get_profile(profile)
    text = report_messages(profile, messages)
    return _create_report(images, output, profile=profile, include_pdf=include_pdf, text=text, prices=prices)


def _create_report(images, output, *, profile, include_pdf, text, prices):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    output, images_dir = Path(output), Path(images)
    data = read_records(output, profile=profile)
    thumbs = output / 'thumbnails'
    thumbs.mkdir(parents=True, exist_ok=True)
    accepted = AcceptedReadings(output)
    for item in data:
        image_path = report_photo(images_dir, item)
        accepted.validate_report(output / 'json' / (Path(item['filename']).stem + '.json'), image_path)
    book = Workbook()

    def header(sheet, labels, widths):
        sheet.append(labels)
        for cell in sheet[1]:
            if isinstance(cell.value, str):
                cell.data_type = 's'
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='60422E')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        for index, label in enumerate(labels, 1):
            sheet.column_dimensions[get_column_letter(index)].width = widths.get(label, 16)

    # Readable upright copies inside the output folder, so links keep working when only this folder is zipped.
    copies = output / 'images'
    copies.mkdir(exist_ok=True)

    def link(cell, name):
        cell.hyperlink = 'images/' + Path(name).stem + '.jpg'
        cell.font = Font(color='0563C1', underline='single')

    literal = book.active
    literal.title = text['sheet.literal']
    literal_keys = ('filename', 'image', 'literal_text', 'handwritten_notes', 'uncertainties', 'reviewed')
    header(literal, [text[f'column.{key}'] for key in literal_keys],
           {text['column.image']:24, text['column.literal_text']:80,
            text['column.handwritten_notes']:32, text['column.uncertainties']:40, text['column.reviewed']:12})
    fields = book.create_sheet(text['sheet.fields'])
    field_keys = [key for key, label in profile.columns if key not in (None, 'literal_text')]
    header(fields, [text[f'field.{key}'] for key in field_keys],
           {text[f'field.{key}']: width for key, width in
            [('locality', 40), ('ecology', 36), ('uncertainties', 40),
             ('name_as_written', 32), ('handwritten_notes', 30)] if key in field_keys})
    for row, item in enumerate(data, 2):
        reading = item['reading']
        doubts = '\n'.join(reading['uncertainties'])
        literal.append([item['filename'], '', profile.literal(reading), reading.get('handwritten_notes', ''), doubts, ''])
        fields.append([item['filename'] if key == 'filename' else doubts if key == 'uncertainties' else json.dumps(reading[key], ensure_ascii=False) if isinstance(reading.get(key), (list, dict)) else reading.get(key, '')
                       for key, _ in profile.columns if key not in (None, 'literal_text')])
        link(literal.cell(row=row, column=1), item['filename'])
        if 'filename' in field_keys:
            link(fields.cell(row=row, column=field_keys.index('filename') + 1), item['filename'])
        image = upright(report_photo(images_dir, item), reading.get('rotation_degrees', 0))
        copy = image.copy()
        copy.thumbnail((2000, 2000), Image.LANCZOS)
        copy.save(copies / (Path(item['filename']).stem + '.jpg'), 'JPEG', quality=82)
        image.thumbnail((170, 170))
        thumb = thumbs / (Path(item['filename']).stem + '.jpg')
        image.save(thumb, 'JPEG', quality=85)
        picture = XLImage(str(thumb))
        literal.add_image(picture, f'B{row}')
        lines = profile.literal(reading).count('\n') + 1
        literal.row_dimensions[row].height = max(picture.height * 0.75 + 6, lines * 15 + 6)
        for sheet in (literal, fields):
            for cell in sheet[row]:
                cell.alignment = Alignment(wrap_text=True, vertical='top')
                # Labels can start with '=' (e.g. '= AMS'); keep them as text, never as an Excel formula.
                if isinstance(cell.value, str):
                    cell.data_type = 's'
    for sheet, frozen in ((literal, 'C2'), (fields, 'B2')):
        sheet.freeze_panes = frozen
        sheet.auto_filter.ref = f'A1:{get_column_letter(sheet.max_column)}{len(data) + 1}'
    xlsx = output / 'transcriptions.xlsx'
    book.save(xlsx)
    pdf = output / 'review.pdf'
    if include_pdf:
        build_pdf(pdf, data, images_dir, profile=profile, _text=text)
    costs = [estimate_cost(item['input_tokens'], item['output_tokens'], prices=prices,
                           model=item.get('model')) for item in data]
    cost = 'unknown (rates not configured for every model)' if any(v is None for v in costs) else f'{sum(costs):.4f} USD'
    print(f"{len(data)} transcriptions: {xlsx.name}{' and ' + pdf.name if include_pdf else ''} · estimated cost: {cost}")
    return ReportFiles(xlsx, pdf if include_pdf else None)


def report(args, *, profile='po-bryophytes', messages=None, prices=None):
    """Compatibility adapter; new integrations use create_report(images, output, ...)."""
    return create_report(args.images, args.output, profile=profile,
                         include_pdf=not args.no_pdf, messages=messages, prices=prices)


def build_pdf(target, data, images_dir, *, profile="po-bryophytes", messages=None, _text=None):
    profile = get_profile(profile)
    text = _text if _text is not None else report_messages(profile, messages)
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import Image as RLImage, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, XPreformatted
    styles = getSampleStyleSheet()
    small = styles['BodyText'].clone('small', fontSize=8.5, leading=10.5)
    story = []
    for index, item in enumerate(data):
        reading = item['reading']
        image = upright(report_photo(images_dir, item), reading.get('rotation_degrees', 0))
        image.thumbnail((1100, 1100))
        buffer = io.BytesIO()
        image.save(buffer, 'JPEG', quality=78)
        buffer.seek(0)
        width, height = ImageReader(buffer).getSize()
        scale = min(180 * mm / width, 105 * mm / height)
        buffer.seek(0)
        story += [Paragraph(f"<b>{escape(item['filename'])}</b> · {index + 1}/{len(data)}", styles['Heading4']),
                  RLImage(buffer, width * scale, height * scale), Spacer(1, 4 * mm)]
        mono = styles['Code'].clone('mono', fontName='Courier', fontSize=9, leading=11.5, backColor=colors.HexColor('#FBF7F1'),
                                    borderColor=colors.HexColor('#C9B8A6'), borderWidth=0.5, borderPadding=6)
        story += [Paragraph(f"<b>{escape(text['column.literal_text'])}</b>", small), Spacer(1, 2 * mm), XPreformatted(escape(profile.literal(reading)), mono), Spacer(1, 5 * mm)]
        extras = [(text['column.handwritten_notes'], reading.get('handwritten_notes', '')),
                  (text['column.uncertainties'], '\n'.join(reading['uncertainties'])),
                  (text['column.reviewed'], text['review.choices'])]
        table = Table([[Paragraph(f'<b>{escape(label)}</b>', small), Paragraph(escape(value or '—').replace('\n', '<br/>'), small)]
                       for label, value in extras], colWidths=[38 * mm, 142 * mm])
        table.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), 0.3, colors.HexColor('#C9B8A6')),
                                   ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#F4ECE0')), ('VALIGN', (0, 0), (-1, -1), 'TOP')]))
        story += [table]
        if index < len(data) - 1:
            story.append(PageBreak())
    SimpleDocTemplate(str(target), pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
                      title=text['pdf.title']).build(story)


def batch_folder(images, output=None, *, profile='po-bryophytes', options=None, client=None,
                 include_pdf=False, messages=None, prices=None, accept_existing=False,
                 request_context=None):
    """Return BatchResult after a complete transcription and checked delivery ZIP.

    Price precedence is prices > options.prices > environment rates, applied
    to both transcription and reports. The supplied options remain unchanged.
    """
    return _batch_folder(images, output, profile=profile, options=options, client=client,
                         include_pdf=include_pdf, messages=messages, prices=prices,
                         accept_existing=accept_existing, request_context=request_context)


def _batch_folder(images, output=None, *, profile='po-bryophytes', options=None, client=None,
                  include_pdf=False, messages=None, prices=None, accept_existing=False,
                  request_context=None, configuration=None):
    profile, options = get_profile(profile), options or TranscriptionOptions()
    if prices is not None:
        options = replace(options, prices=prices)
    text = report_messages(profile, messages)  # Validate and freeze presentation before any paid request.
    images = Path(images)
    output = Path(output) if output is not None else images / 'transcription'
    archive = images / 'transcriptions.zip' if output == images / 'transcription' else output.with_name(output.name + '.zip')
    with output_lock(output):
        totals = _transcribe(images, output, profile, options, client, configuration,
                             accept_existing=accept_existing, request_context=request_context)
        missing = [p.name for p in photos(images) if not (output / 'json' / (p.stem + '.json')).exists()]
        if missing:
            raise _IncompleteBatchError(f'{len(missing)} photographs have no reading (for example {missing[0]}). Run again to retry missing readings only.')
        reports = _create_report(images, output, profile=profile, include_pdf=include_pdf,
                                 text=text, prices=prices if prices is not None else options.prices)
        files = [reports.workbook] + sorted((output / 'images').glob('*.jpg')) + ([reports.pdf] if reports.pdf else [])
        partial = archive.with_suffix('.zip.partial')
        with zipfile.ZipFile(partial, 'w', zipfile.ZIP_DEFLATED) as bundle:
            for path in files:
                bundle.write(path, output.name + '/' + path.relative_to(output).as_posix())
        with zipfile.ZipFile(partial) as bundle:
            if bundle.testzip() is not None:
                raise _IncompleteBatchError('ZIP integrity check failed; the previous archive was not replaced.')
        partial.replace(archive)
        print(f'ZIP ready: {archive} ({archive.stat().st_size / 1e6:.1f} MB, {len(files)} files). Extract the ZIP before opening Excel.')
    return BatchResult(totals, reports, archive)


def batch(args, *, profile='po-bryophytes', options=None, client=None, _configuration=None,
          messages=None, prices=None):
    """Compatibility adapter preserving the earlier namespace mutation/CLI exit contract."""
    args.output = str(Path(args.images) / 'transcription')
    args.limit, args.sample, args.no_pdf = None, False, not args.pdf
    try:
        return _batch_folder(args.images, args.output, profile=profile, options=options,
                             client=client, include_pdf=args.pdf, messages=messages, prices=prices,
                             accept_existing=getattr(args, 'accept_existing', False), configuration=_configuration)
    except _IncompleteBatchError as error:
        raise SystemExit(str(error)) from None
