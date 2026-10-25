"""Transcribe bryophyte herbarium labels (PO / LISU) with Claude and deliver Excel + review PDF.

  batch      --images DIR [--pdf]                                (standard run: transcribe new photos, Excel, zip)
  transcribe --images DIR --output DIR [--limit N] [--sample]   (calls the API; one JSON per photo)
  report     --images DIR --output DIR                            (no API: builds the .xlsx and .pdf)

  demo       --output DIR [--pdf]                              (offline synthetic delivery)

Originals are never modified; rotation is applied only to the copies shown in Excel/PDF.
"""
import argparse
import os
from typing import Literal
import base64
import io
import json
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from importlib.metadata import PackageNotFoundError, version
from xml.sax.saxutils import escape

import zipfile

from PIL import Image, ImageOps
try:
    from pillow_heif import register_heif_opener
except ImportError:
    register_heif_opener = None
from pydantic import BaseModel, ConfigDict, model_validator
from .cache import TaskResult, digest, identity, output_lock, reusable
from .legacy import BRYOPHYTE_FIELDS, RECORD_SCHEMA_VERSION, normalize_record, rename_keys

if register_heif_opener is not None:
    register_heif_opener()

MODEL = os.environ.get('TRANSCRIPTION_MODEL', 'claude-haiku-5-5')

def cost_text(tokens_in, tokens_out):
    price_in, price_out = os.getenv('TRANSCRIPTION_PRICE_IN'), os.getenv('TRANSCRIPTION_PRICE_OUT')
    if not price_in or not price_out:
        return 'unknown (rates not configured)'
    try:
        rates = float(price_in), float(price_out)
        import math
        if any(not math.isfinite(v) or v < 0 for v in rates):
            raise ValueError
    except ValueError:
        raise ValueError('Rates must be finite nonnegative numbers.') from None
    return f'{(tokens_in * rates[0] + tokens_out * rates[1]) / 1e6:.4f} USD'


def safe_photo(folder, name):
    root = Path(folder).resolve()
    candidate = (root / name).resolve()
    if Path(name).name != name or not candidate.is_relative_to(root) or not candidate.is_file():
        raise ValueError('Image reference is outside the input directory or does not exist.')
    return candidate



class BryophyteReading(BaseModel):
    model_config = ConfigDict(extra='forbid')

    @model_validator(mode='before')
    @classmethod
    def read_legacy_fields(cls, value):
        return rename_keys(value, BRYOPHYTE_FIELDS)

    rotation_degrees: Literal[0, 90, 180, 270]
    herbarium: str
    collection_series: str
    name_as_written: str
    species: str
    taxonomic_author: str
    determiner: str
    country_province: str
    locality: str
    ecology: str
    altitude: str
    utm: str
    date_as_written: str
    date_iso: str
    collector: str
    collection_number: str
    handwritten_notes: str
    literal_text: str
    uncertainties: list[str]


PROMPT = """This photograph shows a bryophyte herbarium label, usually from "Herbário da Universidade do Porto — Instituto de Botânica Dr. Gonçalo Sampaio (PO)", the "Briófitas das comunidades de Sphagnum" series, or the herbarium of Museu, Laboratório e Jardim Botânico de Lisboa (LISU). The label may be rotated.

Usual fields: scientific name with author; "Det." (determiner); "Stat." (station: a country/province code such as "Lu: TM", "Lu: Mi", or "Be:", followed by locality); "Ecol." (ecology); "Alt."; "UTM"; "Data" (date); "Leg." (collector); "Nº" (collection number).

Rules:
- Copy the text exactly as written. Do not correct names, spelling, or abbreviations, and do not complete unreadable text. Use an empty string for an absent or illegible field. Never invent information. Preserve the source language; do not translate label text into English.
- rotation_degrees: clockwise degrees (0, 90, 180, or 270) needed to make the photograph readable.
- name_as_written: the entire name as it appears, including varieties, forms, and parenthesized synonyms. species: the scientific name only (genus, epithet, and infraspecific designation, without author). taxonomic_author: the author only.
- country_province: the complete country/province code at the start of "Stat." (for example "Lu: TM", "Lu: Mi", "Be:"), even when its colon is absent. locality: the remaining text without that code.
- date_iso: YYYY-MM-DD only when the date is unambiguous. These labels often use YY/MM/DD (for example 89/05/09 = 1989-05-09), but also D/M/YY (for example 15/3/93). If the order is uncertain, leave date_iso empty and explain in uncertainties.
- handwritten_notes: handwritten text, such as name corrections, "Rev.", or signatures, with its position.
- literal_text: the most faithful possible transcription, formatted as follows:
  * first line: an editorial origin marker, [Porto] for PO, [LISU] for Lisbon, or [Other: name];
  * then reproduce each line in label order, including printed field labels (Det., Stat., Ecol., Alt., UTM, Data, Leg., Nº);
  * fields on the same label line remain on the same line, separated by four spaces;
  * place handwriting where it appears, within the editorial marker [handwritten: ...];
  * if the photograph contains several labels, use a separate block for each, separated by a blank line. Do not merge them.
  Example format (retain the original-language text):
  [Porto]
  Herbário da Universidade do Porto
  Instituto de Botânica Dr. Gonçalo Sampaio (PO)
  Briófitas das comunidades de Sphagnum
  Sphagnum palustre L.
  Det. Ana Séneca
  Stat. Lu: TM, Vila Real, Lamas de Olo, PNSA, antes da ponte da aldeia
  Ecol. com Molinia
  Alt. 1000 m    UTM PF0180    Data 15/3/93
  Leg. Ana Séneca    Nº 565
- uncertainties: genuinely uncertain readings only, identifying the field, proposed reading, and reason. Do not list confirmations or absent fields."""

COLUMNS = [('filename', 'Filename'), (None, 'Image'), ('herbarium', 'Herbarium'), ('collection_series', 'Collection / series'),
           ('name_as_written', 'Name as written'), ('species', 'Species'), ('taxonomic_author', 'Taxonomic author'),
           ('determiner', 'Determiner'), ('country_province', 'Country / province'), ('locality', 'Locality (Stat.)'),
           ('ecology', 'Ecology'), ('altitude', 'Altitude'), ('utm', 'UTM'), ('date_as_written', 'Date as written'),
           ('date_iso', 'ISO date'), ('collector', 'Collector (Leg.)'), ('collection_number', 'Nº'), ('handwritten_notes', 'Handwritten notes'),
           ('uncertainties', 'Reading uncertainties'), ('literal_text', 'Literal text'), (None, 'Reviewed')]


def upright(path, rotation=0):
    image = ImageOps.exif_transpose(Image.open(path)).convert('RGB')
    return image.rotate(-rotation, expand=True) if rotation in (90, 180, 270) else image


def encoded(image, long_edge=1568):
    image = image.copy()
    image.thumbnail((long_edge, long_edge), Image.LANCZOS)
    buffer = io.BytesIO()
    image.save(buffer, 'JPEG', quality=92)
    return {'type':'image', 'source':{'type':'base64', 'media_type':'image/jpeg', 'data':base64.standard_b64encode(buffer.getvalue()).decode()}}


def photos(folder):
    files = sorted(p for p in Path(folder).iterdir() if p.is_file() and p.suffix.lower() in ('.jpg', '.jpeg', '.png', '.heic', '.heif'))
    stems = [p.stem.casefold() for p in files]
    if len(stems) != len(set(stems)):
        raise ValueError('Duplicate image filename stems: rename the images before processing.')
    for path in files:
        safe_photo(folder, path.name)
    return files


def cache_configuration():
    try:
        provider_version = version('anthropic')
    except PackageNotFoundError:
        provider_version = None  # Offline tests may provide a simulated client only.
    return {'model': MODEL, 'prompt': PROMPT, 'schema': BryophyteReading.model_json_schema(),
            'anthropic': provider_version, 'pydantic': version('pydantic'),
            'max_tokens': 16000, 'transform': 'exif-rgb-thumbnail-jpeg-v1',
            'long_edge': 1568, 'jpeg_quality': 92, 'pillow': version('Pillow'),
            'pillow_heif': version('pillow-heif') if register_heif_opener else None}


def transcribe(args):
    with output_lock(args.output):
        return _transcribe(args)


def _transcribe(args):
    import anthropic
    key = os.environ.get('ANTHROPIC_API_KEY', '').strip()
    if not key:
        raise ValueError('Set ANTHROPIC_API_KEY in the environment.')
    out = Path(args.output) / 'json'
    out.mkdir(parents=True, exist_ok=True)
    files = photos(args.images)
    if args.sample and args.limit:
        step = max(1, len(files) // args.limit)
        files = files[::step][:args.limit]
    elif args.limit:
        files = files[:args.limit]
    # Check the entire selected batch before making any paid request.
    configuration = cache_configuration()
    tasks = []
    for path in files:
        target = out / (path.stem + '.json')
        expected = identity(path, configuration)
        cached = reusable(target, path.name, expected, BryophyteReading)
        tasks.append((path, target, expected, cached))
    client = anthropic.Anthropic(api_key=key, timeout=120, max_retries=3)

    def run(task):
        path, target, expected, cached = task
        if cached:
            return TaskResult(reused=1)
        content = path.read_bytes()
        if digest(content) != expected['image_sha256']:
            raise ValueError(f'{path.name}: image changed during the run; run the command again.')
        started = time.time()
        try:
            response = client.messages.parse(model=MODEL, max_tokens=16000, output_format=BryophyteReading, messages=[
                {'role':'user', 'content':[encoded(upright(io.BytesIO(content))), {'type':'text', 'text':PROMPT}]}])
        except anthropic.APIError as error:
            print(f'{path.name}: API error ({type(error).__name__})')
            return TaskResult(failures=1)
        if response.stop_reason != 'end_turn' or response.parsed_output is None:
            print(f'{path.name}: incomplete reading ({response.stop_reason})')
            return TaskResult(failures=1, input_tokens=response.usage.input_tokens,
                              output_tokens=response.usage.output_tokens)
        reading = BryophyteReading.model_validate(response.parsed_output)
        partial = target.with_suffix('.json.partial')
        partial.write_text(json.dumps({'schema_version':RECORD_SCHEMA_VERSION, 'filename':path.name, 'model':MODEL, 'elapsed_seconds':round(time.time() - started, 1),
                                      'input_tokens':response.usage.input_tokens, 'output_tokens':response.usage.output_tokens,
                                      'cache':expected, 'reading':reading.model_dump()}, ensure_ascii=False, indent=2), encoding='utf-8')
        if target.exists():
            raise ValueError(f'{target.name}: a result appeared during the run; it was not replaced.')
        partial.replace(target)
        return TaskResult(photos=1, input_tokens=response.usage.input_tokens,
                          output_tokens=response.usage.output_tokens)

    with ThreadPoolExecutor(max_workers=6) as pool:
        results = list(pool.map(run, tasks))
    totals = TaskResult(**{name: sum(getattr(result, name) for result in results)
                           for name in TaskResult.__dataclass_fields__})
    cost = cost_text(totals.input_tokens, totals.output_tokens)
    print(f'transcribed {totals.photos} · reused {totals.reused} · failures {totals.failures} '
          f'· tokens {totals.input_tokens}+{totals.output_tokens} · estimated cost: {cost}')
    return totals


def read_records(output):
    results = [normalize_record(json.loads(p.read_text(encoding='utf-8'))) for p in sorted((Path(output) / 'json').glob('*.json'))]
    for item in results:
        BryophyteReading.model_validate(item['reading'])
    return results


def report_photo(folder, item):
    path = safe_photo(folder, item['filename'])
    if 'cache' in item and digest(path.read_bytes()) != item['cache']['image_sha256']:
        raise ValueError(f"{path.name}: image differs from the transcribed source; restore it or use a new output directory.")
    return path


def report(args):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    output, images_dir = Path(args.output), Path(args.images)
    thumbs = output / 'thumbnails'
    thumbs.mkdir(parents=True, exist_ok=True)
    data = read_records(output)
    for item in data:
        report_photo(images_dir, item)
    book = Workbook()

    def header(sheet, labels, widths):
        sheet.append(labels)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='60422E')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        for index, label in enumerate(labels, 1):
            sheet.column_dimensions[get_column_letter(index)].width = widths.get(label, 16)

    # Readable upright copies inside the output folder, so links keep working when only this folder is zipped.
    copies = output / 'images'
    copies.mkdir(exist_ok=True)

    def link(cell, name):
        cell.hyperlink = 'images/' + Path(name).stem + '.jpg'
        cell.font = Font(color='0563C1', underline='single')

    literal = book.active
    literal.title = 'Literal transcription'
    header(literal, ['Filename', 'Image', 'Literal transcription', 'Handwritten notes', 'Reading uncertainties', 'Reviewed'],
           {'Image':24, 'Literal transcription':80, 'Handwritten notes':32, 'Reading uncertainties':40, 'Reviewed':12})
    fields = book.create_sheet('Fields')
    header(fields, [label for key, label in COLUMNS if key not in (None, 'literal_text')],
           {'Locality (Stat.)':40, 'Ecology':36, 'Reading uncertainties':40, 'Name as written':32, 'Handwritten notes':30})
    for row, item in enumerate(data, 2):
        reading = item['reading']
        doubts = '\n'.join(reading['uncertainties'])
        literal.append([item['filename'], '', reading['literal_text'], reading['handwritten_notes'], doubts, ''])
        fields.append([item['filename'] if key == 'filename' else doubts if key == 'uncertainties' else reading.get(key, '')
                       for key, _ in COLUMNS if key not in (None, 'literal_text')])
        link(literal.cell(row=row, column=1), item['filename'])
        link(fields.cell(row=row, column=1), item['filename'])
        image = upright(report_photo(images_dir, item), reading['rotation_degrees'])
        copy = image.copy()
        copy.thumbnail((2000, 2000), Image.LANCZOS)
        copy.save(copies / (Path(item['filename']).stem + '.jpg'), 'JPEG', quality=82)
        image.thumbnail((170, 170))
        thumb = thumbs / (Path(item['filename']).stem + '.jpg')
        image.save(thumb, 'JPEG', quality=85)
        picture = XLImage(str(thumb))
        literal.add_image(picture, f'B{row}')
        lines = reading['literal_text'].count('\n') + 1
        literal.row_dimensions[row].height = max(picture.height * 0.75 + 6, lines * 15 + 6)
        for sheet in (literal, fields):
            for cell in sheet[row]:
                cell.alignment = Alignment(wrap_text=True, vertical='top')
                # Labels can start with '=' (e.g. '= AMS'); keep them as text, never as an Excel formula.
                if isinstance(cell.value, str):
                    cell.data_type = 's'
    for sheet, frozen in ((literal, 'C2'), (fields, 'B2')):
        sheet.freeze_panes = frozen
        sheet.auto_filter.ref = f'A1:{get_column_letter(sheet.max_column)}{len(data) + 1}'
    xlsx = output / 'transcriptions.xlsx'
    book.save(xlsx)
    pdf = output / 'review.pdf'
    if not args.no_pdf:
        build_pdf(pdf, data, images_dir)
    tokens_in, tokens_out = sum(i['input_tokens'] for i in data), sum(i['output_tokens'] for i in data)
    print(f"{len(data)} transcriptions: {xlsx.name}{'' if args.no_pdf else ' and ' + pdf.name} · estimated cost: {cost_text(tokens_in, tokens_out)}")


def build_pdf(target, data, images_dir):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import Image as RLImage, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, XPreformatted
    styles = getSampleStyleSheet()
    small = styles['BodyText'].clone('small', fontSize=8.5, leading=10.5)
    story = []
    for index, item in enumerate(data):
        reading = item['reading']
        image = upright(report_photo(images_dir, item), reading['rotation_degrees'])
        image.thumbnail((1100, 1100))
        buffer = io.BytesIO()
        image.save(buffer, 'JPEG', quality=78)
        buffer.seek(0)
        width, height = ImageReader(buffer).getSize()
        scale = min(180 * mm / width, 105 * mm / height)
        buffer.seek(0)
        story += [Paragraph(f"<b>{escape(item['filename'])}</b> · {index + 1}/{len(data)}", styles['Heading4']),
                  RLImage(buffer, width * scale, height * scale), Spacer(1, 4 * mm)]
        mono = styles['Code'].clone('mono', fontName='Courier', fontSize=9, leading=11.5, backColor=colors.HexColor('#FBF7F1'),
                                    borderColor=colors.HexColor('#C9B8A6'), borderWidth=0.5, borderPadding=6)
        story += [Paragraph('<b>Literal transcription</b>', small), Spacer(1, 2 * mm), XPreformatted(escape(reading['literal_text']), mono), Spacer(1, 5 * mm)]
        extras = [('Handwritten notes', reading['handwritten_notes']), ('Reading uncertainties', '\n'.join(reading['uncertainties'])),
                  ('Reviewed', '[ ] yes    [ ] corrected')]
        table = Table([[Paragraph(f'<b>{escape(label)}</b>', small), Paragraph(escape(value or '—').replace('\n', '<br/>'), small)]
                       for label, value in extras], colWidths=[38 * mm, 142 * mm])
        table.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), 0.3, colors.HexColor('#C9B8A6')),
                                   ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#F4ECE0')), ('VALIGN', (0, 0), (-1, -1), 'TOP')]))
        story += [table]
        if index < len(data) - 1:
            story.append(PageBreak())
    SimpleDocTemplate(str(target), pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
                      title='Bryophyte transcription review').build(story)


def batch(args):
    """Standard delivery, with exclusive ownership of the output during the run."""
    args.output = str(Path(args.images) / 'transcription')
    with output_lock(args.output):
        return _batch(args)


def _batch(args):
    images = Path(args.images)
    args.output = str(images / 'transcription')
    args.limit, args.sample, args.no_pdf = None, False, not args.pdf
    pending = [p for p in photos(images) if not (Path(args.output) / 'json' / (p.stem + '.json')).exists()]
    print(f'{len(photos(images))} photographs · {len(pending)} pending; cost depends on tokens and rates.')
    _transcribe(args)
    missing = [p.name for p in photos(images) if not (Path(args.output) / 'json' / (p.stem + '.json')).exists()]
    if missing:
        sys.exit(f'{len(missing)} photographs have no reading (for example {missing[0]}). Run again to retry missing readings only.')
    report(args)
    archive = images / 'transcriptions.zip'
    output = Path(args.output)
    files = [output / 'transcriptions.xlsx'] + sorted((output / 'images').glob('*.jpg')) + ([output / 'review.pdf'] if args.pdf else [])
    partial = archive.with_suffix('.zip.partial')
    with zipfile.ZipFile(partial, 'w', zipfile.ZIP_DEFLATED) as bundle:
        for path in files:
            bundle.write(path, 'transcription/' + path.relative_to(output).as_posix())
    with zipfile.ZipFile(partial) as bundle:
        if bundle.testzip() is not None:
            sys.exit('ZIP integrity check failed; the previous archive was not replaced.')
    partial.replace(archive)
    print(f'ZIP ready: {archive} ({archive.stat().st_size / 1e6:.1f} MB, {len(files)} files). Extract the ZIP before opening Excel.')


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    commands = parser.add_subparsers(dest='command', required=True)
    demo_command = commands.add_parser('demo', help='Create a synthetic delivery without network or API usage.')
    demo_command.add_argument('--output', required=True, help='New directory for demonstration files.')
    demo_command.add_argument('--pdf', action='store_true', help='Also generate a review PDF.')
    for name in ('transcribe', 'report', 'batch'):
        command = commands.add_parser(name)
        command.add_argument('--images', required=True)
        if name == 'batch':
            command.add_argument('--pdf', action='store_true', help='Also generate the review PDF.')
            continue
        command.add_argument('--output', required=True)
        if name == 'report':
            command.add_argument('--no-pdf', action='store_true')
        if name == 'transcribe':
            command.add_argument('--limit', type=int)
            command.add_argument('--sample', action='store_true')
    args = parser.parse_args()
    if args.command == 'demo':
        from .demo import run_demo
        run_demo(args.output, include_pdf=args.pdf)
        return
    if not Path(args.images).is_dir():
        sys.exit(f'Image directory not found: {args.images}')
    {'transcribe':transcribe, 'report':report, 'batch':batch}[args.command](args)


if __name__ == '__main__':
    main()
