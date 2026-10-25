"""Generate and verify a synthetic delivery with a simulated provider, offline."""
import hashlib
import json
import os
from pathlib import Path, PurePosixPath
import sys
from types import SimpleNamespace
from unittest.mock import patch
import zipfile

from PIL import Image, ImageDraw, ImageFont

from . import bryophytes, sennen


def run_demo(destination, include_pdf=False):
    """Create a new demo directory; never overwrite an existing one or call a provider."""
    from openpyxl import load_workbook

    destination = Path(destination)
    destination.mkdir(parents=True, exist_ok=False)
    images = destination / 'input'
    images.mkdir()
    image_path = images / 'synthetic.jpg'
    literal = '= SYNTHETIC LABEL - NOT A SPECIMEN\nSphagnum palustre L.\nLeg. Example collector    No. 001\nDate 15/3/1993'
    image = Image.new('RGB', (1000, 700), '#faf7ed')
    drawing = ImageDraw.Draw(image)
    drawing.rectangle((30, 30, 970, 670), outline='#60422e', width=4)
    drawing.multiline_text((60, 90), literal, font=ImageFont.load_default(size=26),
                           fill='#201a16', spacing=18)
    image.save(image_path, 'JPEG', quality=92)
    original_hash = hashlib.sha256(image_path.read_bytes()).hexdigest()
    fields = {name: '' for name in bryophytes.BryophyteReading.model_fields}
    fields.update(rotation_degrees=0, herbarium='Synthetic collection',
                  species='Sphagnum palustre', taxonomic_author='L.',
                  collector='Example collector', collection_number='001',
                  date_as_written='15/3/1993', date_iso='1993-03-15',
                  literal_text=literal, uncertainties=['Synthetic uncertainty for review.'])
    reading = bryophytes.BryophyteReading(**fields)
    sheet_fields = {name: '' for name in sennen.SheetReading.model_fields}
    sheet_fields.update(labels=[sennen.Label(origin='Sennen', literal_text='SYNTHETIC LABEL A'),
                                sennen.Label(origin='Porto', literal_text='SYNTHETIC LABEL B')],
                        uncertainties=[])
    sheet = sennen.SheetReading(**sheet_fields)
    calls = []

    def parse(**kwargs):
        is_sheet = kwargs['output_format'] is sennen.SheetReading
        calls.append('sennen' if is_sheet else 'bryophytes')
        return SimpleNamespace(stop_reason='end_turn', parsed_output=sheet if is_sheet else reading,
                               usage=SimpleNamespace(input_tokens=10, output_tokens=20))

    class SimulatedAPIError(Exception):
        pass

    client = SimpleNamespace(messages=SimpleNamespace(parse=parse))
    provider = SimpleNamespace(APIError=SimulatedAPIError, Anthropic=lambda **kwargs: client)
    args = SimpleNamespace(images=images, pdf=include_pdf)
    settings = {'ANTHROPIC_API_KEY': 'offline-demo-placeholder',
                'TRANSCRIPTION_PRICE_IN': '', 'TRANSCRIPTION_PRICE_OUT': ''}
    with patch.dict(sys.modules, {'anthropic': provider}), patch.dict(os.environ, settings), \
            patch.object(bryophytes, 'MODEL', 'offline-demo'):
        bryophytes.batch(args)
        resumed = bryophytes.transcribe(args)
        sheet_reading, sheet_usage = sennen.read_labels(image_path.read_bytes(), client=client)
    (destination / 'sheet_reading.json').write_text(
        json.dumps({'reading': sheet_reading.model_dump(), 'usage': sheet_usage},
                   ensure_ascii=False, indent=2), encoding='utf-8')
    output = images / 'transcription'
    record = json.loads((output / 'json/synthetic.json').read_text(encoding='utf-8'))
    book = load_workbook(output / 'transcriptions.xlsx')
    try:
        english_sheets = book.sheetnames == ['Literal transcription', 'Fields']
        formula_as_text = book.worksheets[0]['C2'].data_type == 's'
        link = book.worksheets[0]['A2'].hyperlink.target
        relative_link = (output / link).is_file()
    finally:
        book.close()
    extracted = destination / 'extracted'
    with zipfile.ZipFile(images / 'transcriptions.zip') as archive:
        zip_valid = archive.testzip() is None
        safe_names = all(not PurePosixPath(name).is_absolute() and '..' not in PurePosixPath(name).parts
                         for name in archive.namelist())
        if not safe_names:
            raise ValueError('Unexpected unsafe path in generated archive.')
        archive.extractall(extracted)
    checks = {
        'original_unchanged': hashlib.sha256(image_path.read_bytes()).hexdigest() == original_hash,
        'english_json': 'reading' in record and 'literal_text' in record['reading'] and 'leitura' not in record,
        'literal_preserved': record['reading']['literal_text'] == literal,
        'english_workbook_sheets': english_sheets,
        'formula_like_text_is_safe': formula_as_text,
        'image_link_exists': relative_link,
        'extracted_image_link_exists': (extracted / 'transcription' / link).is_file(),
        'zip_valid': zip_valid and safe_names,
        'cache_reused_without_call': resumed.reused == 1 and calls.count('bryophytes') == 1,
        'sheet_labels_separate': len(sheet_reading.labels) == 2,
    }
    if include_pdf:
        checks['pdf_generated'] = (output / 'review.pdf').read_bytes().startswith(b'%PDF')
    result = {'provider': 'simulated; no network or paid requests', 'checks': checks,
              'all_passed': all(checks.values()), 'simulated_calls': len(calls)}
    (destination / 'verification.json').write_text(json.dumps(result, indent=2), encoding='utf-8')
    if not result['all_passed']:
        raise RuntimeError('Offline demonstration verification failed; inspect verification.json.')
    print(f'Offline demo verified: {destination / "verification.json"}')
    return result
